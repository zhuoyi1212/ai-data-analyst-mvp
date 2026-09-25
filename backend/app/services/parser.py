"""上传文件解析与边界校验（FR-1）。

只做「读成标准二维表」，不做语义推断（语义是 Semantic Profiler 的职责）。
"""
from __future__ import annotations

import csv
import io
import math
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from app.config import settings


class UploadError(ValueError):
    """文件不满足上传边界，message 为可直接展示的中文原因。"""


ALLOWED_SUFFIX = {".csv", ".xlsx", ".xls"}
_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk")


@dataclass
class ParsedTable:
    filename: str
    df: pd.DataFrame
    physical_types: dict[str, str] = field(default_factory=dict)

    @property
    def rows(self) -> int:
        return int(self.df.shape[0])

    @property
    def cols(self) -> int:
        return int(self.df.shape[1])

    def preview(self, n: int = settings.preview_rows) -> list[dict]:
        return [
            {col: json_safe(val) for col, val in row.items()}
            for row in self.df.head(n).to_dict(orient="records")
        ]


def json_safe(value):
    """把 pandas/numpy 标量转成可 JSON 序列化的 Python 原生值。"""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if pd.isna(value) if not isinstance(value, (list, dict)) else False:
        return None
    return value


def friendly_dtype(series: pd.Series) -> str:
    dtype = str(series.dtype)
    if dtype.startswith(("int", "Int", "uint")):
        return "integer"
    if dtype.startswith(("float", "Float")):
        return "float"
    if dtype.startswith("bool"):
        return "boolean"
    if dtype.startswith("datetime"):
        return "datetime"
    if dtype.startswith("date"):
        return "date"
    return "string"


def normalize_mixed_columns(df: pd.DataFrame) -> pd.DataFrame:
    """把 object/混合类型列归一化为明确 dtype（写 Parquet 前必需）。

    背景：Excel 常把同一列读成 object 混合类型（如 Postal Code 列里美国邮编是
    整数、加拿大邮编是 'M7A' 文本），pyarrow 无法为该列推断统一 schema，
    to_parquet 会抛 "Could not convert 'M7A' ... tried to convert to int64"。

    规则：
    - 非 object 列保持不变（已由质量处理转成数值/日期的列不受影响）；
    - object 列：除空值外全部可解析为数值 → 转 float64，保留数值语义；
      否则转 pandas string，按文本处理（邮编、订单号等本就该是文本）。
    数值计算不受影响：引擎对指标列统一 pd.to_numeric(errors="coerce")。
    """
    out = df.copy()
    for col in out.columns:
        series = out[col]
        if str(series.dtype) != "object":
            continue
        non_null = series.dropna()
        if non_null.empty:
            continue
        # 纯 Python 文本列 pyarrow 可直接推断为 string，保持原样（避免无谓的字节变化）；
        # 只有混入 int/float/bool 等标量（数字+文本混合）时才需归一化。
        python_types = {type(v) for v in non_null}
        if python_types <= {str}:
            continue
        converted = pd.to_numeric(non_null, errors="coerce")
        if converted.notna().all():
            out[col] = pd.to_numeric(series, errors="coerce")
        else:
            out[col] = series.map(lambda v: str(v) if pd.notna(v) else None).astype("string")
    return out


def _read_csv_bytes(data: bytes) -> pd.DataFrame:
    text = None
    for enc in _ENCODINGS:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise UploadError("文件编码无法识别，请将 CSV 另存为 UTF-8 编码后重试。")

    sample = text[:4096]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass

    return pd.read_csv(io.StringIO(text), sep=delimiter, dtype=None)


def _read_excel_bytes(data: bytes, suffix: str) -> pd.DataFrame:
    engine = "openpyxl" if suffix == ".xlsx" else "xlrd"
    try:
        return pd.read_excel(io.BytesIO(data), sheet_name=0, engine=engine)
    except Exception as exc:  # 解析失败统一为可展示错误
        raise UploadError(f"Excel 解析失败（仅支持首个工作表的单一二维表）：{exc}") from exc


def _raw_header(data: bytes, suffix: str) -> list[str]:
    """读取未经 pandas 改写的原始表头（pandas 会把重名列改成 a.1）。"""
    if suffix == ".csv":
        text = None
        for enc in _ENCODINGS:
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise UploadError("文件编码无法识别，请将 CSV 另存为 UTF-8 编码后重试。")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return []
        sample = text[:4096]
        delimiter = ","
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            pass
        return next(csv.reader(io.StringIO(lines[0]), delimiter=delimiter))

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        wb.close()
        return ["" if v is None else str(v) for v in row]

    import xlrd

    wb = xlrd.open_workbook(file_contents=data)
    return [str(v) for v in wb.sheet_by_index(0).row_values(0)]


def parse_table(
    filename: str,
    data: bytes,
    *,
    max_file_mb: int | None = None,
    max_rows: int | None = None,
    max_cols: int | None = None,
) -> ParsedTable:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIX:
        raise UploadError("不支持的文件类型：仅支持 .csv / .xlsx / .xls 文件。")

    max_file_mb = settings.max_file_mb if max_file_mb is None else max_file_mb
    max_rows = settings.max_rows if max_rows is None else max_rows
    max_cols = settings.max_cols if max_cols is None else max_cols

    size_mb = len(data) / (1024 * 1024)
    if size_mb > max_file_mb:
        raise UploadError(f"文件过大：{size_mb:.1f}MB，单文件上限为 {max_file_mb}MB。")
    if len(data) == 0:
        raise UploadError("文件为空，请上传包含数据的文件。")

    # 先对原始表头做校验（pandas 会自动改写空列名/重名列，掩盖问题）
    raw_headers = [h.strip() for h in _raw_header(data, suffix)]
    if not raw_headers:
        raise UploadError("未检测到有效表头，请确认首行为字段名称。")
    if any(not h for h in raw_headers):
        raise UploadError("检测到空表头列，请补全所有列名后重试。")
    if len(set(raw_headers)) != len(raw_headers):
        raise UploadError("表头存在重复列名，请保证每列名称唯一。")

    try:
        df = _read_csv_bytes(data) if suffix == ".csv" else _read_excel_bytes(data, suffix)
    except UploadError:
        raise
    except Exception as exc:
        raise UploadError(f"文件解析失败，请确认是标准二维表格：{exc}") from exc

    df.columns = [str(c).strip() for c in df.columns]

    if df.shape[0] == 0:
        raise UploadError("未检测到数据行，文件仅包含表头。")
    if df.shape[0] > max_rows:
        raise UploadError(f"数据行数 {df.shape[0]:,} 超出上限 {max_rows:,} 行。")
    if df.shape[1] > max_cols:
        raise UploadError(f"列数 {df.shape[1]} 超出上限 {max_cols} 列。")

    physical_types = {col: friendly_dtype(df[col]) for col in df.columns}
    return ParsedTable(filename=filename, df=df, physical_types=physical_types)
